"""Reads an examiner's CSV / XLSX / PDF of multiple-choice questions.

Every row must have ALL of: subject, difficulty, topic, question, option_a..option_d,
correct_option (A-D) and a solution. Anything missing or invalid is reported as an error
with its row/line number; the import route then refuses the whole file (nothing is saved).

CSV / XLSX / PDF-table header (any order, these spellings are accepted):
  Subject | Difficulty | Topic | Question | Option A | Option B | Option C | Option D | Correct Answer | Solution

PDF can also be written as labelled blocks:
  Subject: Physics
  Difficulty: Easy
  Topic: Kinematics
  Question: What is the SI unit of velocity?
  A) m/s
  B) km
  C) kg
  D) s        (one option per line)
  Answer: A
  Solution: Velocity is distance per unit time, so m/s.
(every field is required for every question)
"""
import csv
import io
import re
from dataclasses import dataclass, field

REQUIRED = ("subject", "difficulty", "topic", "question",
            "option_a", "option_b", "option_c", "option_d", "correct_option", "solution")
FIELDS = REQUIRED
OPTION_KEYS = ("option_a", "option_b", "option_c", "option_d")

# Column-header spellings we accept (lower-case, underscores/extra spaces removed) -> field name
HEADER_ALIASES = {
    "subject": "subject", "course": "subject",
    "difficulty": "difficulty", "level": "difficulty", "difficulty level": "difficulty",
    "topic": "topic", "chapter": "topic",
    "question": "question", "question text": "question", "questions": "question",
    "solution": "solution", "explanation": "solution", "answer explanation": "solution",
    "working": "solution", "solution explanation": "solution",
    "correct": "correct_option", "correct answer": "correct_option", "correct option": "correct_option",
    "answer": "correct_option", "ans": "correct_option", "right answer": "correct_option",
}
for _letter in "abcd":
    for _spelling in (f"option {_letter}", f"option{_letter}", f"opt {_letter}", f"choice {_letter}", _letter):
        HEADER_ALIASES[_spelling] = f"option_{_letter}"

DIFFICULTY_ALIASES = {
    "easy": "easy", "e": "easy", "beginner": "easy", "basic": "easy", "low": "easy",
    "medium": "medium", "m": "medium", "moderate": "medium", "intermediate": "medium", "med": "medium",
    "hard": "hard", "h": "hard", "difficult": "hard", "advanced": "hard", "high": "hard",
}


@dataclass
class ParseResult:
    rows: list = field(default_factory=list)    # list[dict]
    errors: list = field(default_factory=list)  # list[{"row": int, "reason": str}]


def _clean(value) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def _header_name(cell) -> str:
    return re.sub(r"\s+", " ", _clean(cell).lower().replace("_", " ").replace("-", " ")).strip()


def _resolve_correct(raw: str, options: list):
    """Turns 'B', '(b)', 'Option B', '2' or the exact option text into 'A'..'D'. None if it can't."""
    v = _clean(raw)
    if not v:
        return None
    t = re.sub(r"^(option|opt|choice|ans(wer)?)\s*[:\-]?\s*", "", v, flags=re.I)
    t = t.strip().strip("().[]:").strip()
    if len(t) == 1 and t.upper() in "ABCD":
        return t.upper()
    if t in ("1", "2", "3", "4"):
        return "ABCD"[int(t) - 1]
    for letter, text in zip("ABCD", options):
        if _clean(text).lower() == v.lower():
            return letter
    return None


def _normalize_row(raw: dict, row_no: int, result: ParseResult) -> None:
    row = {k: _clean(raw.get(k)) for k in FIELDS}
    if not any(row.values()):
        return  # blank line, ignore silently
    missing = [k.replace("_", " ") for k in REQUIRED if not row[k]]
    if missing:
        result.errors.append({"row": row_no, "reason": f"Missing {', '.join(missing)}"})
        return
    diff = DIFFICULTY_ALIASES.get(row["difficulty"].lower())
    if not diff:
        result.errors.append({
            "row": row_no,
            "reason": f"Difficulty '{row['difficulty']}' is not easy, medium or hard",
        })
        return
    row["difficulty"] = diff
    options = [row[k] for k in OPTION_KEYS]
    if len({o.lower() for o in options}) < 4:
        result.errors.append({"row": row_no, "reason": "The four options must all be different"})
        return
    correct = _resolve_correct(row["correct_option"], options)
    if not correct:
        result.errors.append({
            "row": row_no,
            "reason": f"Correct answer '{row['correct_option']}' must be A, B, C or D",
        })
        return
    row["correct_option"] = correct
    result.rows.append(row)


def _map_headers(header_cells):
    """Returns {column_index: field} if every required field is present, else None."""
    mapping = {}
    for i, cell in enumerate(header_cells):
        name = HEADER_ALIASES.get(_header_name(cell))
        if name and name not in mapping.values():
            mapping[i] = name
    return mapping if set(REQUIRED) <= set(mapping.values()) else None


HEADER_HELP = ("The header row must contain all of these columns: Subject, Difficulty, Topic, Question, "
               "Option A, Option B, Option C, Option D, Correct Answer, Solution")


def _rows_from_grid(grid, result: ParseResult) -> None:
    grid = [r for r in grid if r is not None]
    if not grid:
        return
    mapping = _map_headers(grid[0])
    if mapping is None:
        result.errors.append({"row": 1, "reason": HEADER_HELP})
        return
    for row_no, cells in enumerate(grid[1:], start=2):
        raw = {f: (cells[i] if i < len(cells) else "") for i, f in mapping.items()}
        _normalize_row(raw, row_no, result)


# ---------- CSV ----------
def parse_csv(data: bytes) -> ParseResult:
    result = ParseResult()
    text = data.decode("utf-8-sig", errors="replace")
    first = text.split("\n", 1)[0]
    delimiter = max(",;\t", key=first.count)          # choose from the header line
    _rows_from_grid(list(csv.reader(io.StringIO(text), delimiter=delimiter)), result)
    return result


# ---------- XLSX ----------
def parse_xlsx(data: bytes) -> ParseResult:
    from openpyxl import load_workbook
    result = ParseResult()
    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    _rows_from_grid([list(r) for r in ws.iter_rows(values_only=True)], result)
    return result


# ---------- PDF ----------
LABEL_RE = re.compile(
    r"^\s*(subject|difficulty|level|topic|chapter|question|"
    r"option\s*[a-d]|[a-d]|"
    r"correct(?:\s*(?:answer|option))?|answer|"
    r"solution|explanation)\s*[:\-\.\)]\s*(.*)$", re.I)
PAGE_NOISE_RE = re.compile(r"^\s*(page\s*)?\d+(\s*(of|/)\s*\d+)?\s*$", re.I)
TEXT_KEYS = ("question", "solution") + OPTION_KEYS      # may wrap onto following lines


def _label_key(label: str) -> str:
    k = re.sub(r"\s+", " ", label.lower()).strip()
    if k in ("level",):
        return "difficulty"
    if k == "chapter":
        return "topic"
    if k.startswith("correct") or k == "answer":
        return "correct_option"
    if k == "explanation":
        return "solution"
    if k.startswith("option"):
        return "option_" + k[-1]
    if len(k) == 1 and k in "abcd":
        return "option_" + k
    return k


def parse_pdf(data: bytes) -> ParseResult:
    """Handles two layouts:
    1. A table with the columns above.
    2. Labelled blocks (Subject: ... / Difficulty: ... / Question: ... / A) ... / Answer: B / Solution: ...)
    """
    import pdfplumber
    result = ParseResult()
    grid, lines = [], []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                grid.extend([_clean(c) for c in r] for r in table)
            lines.extend((page.extract_text() or "").splitlines())

    # Layout 1: table. Repeated header rows on later pages are skipped.
    header_idx = next((i for i, r in enumerate(grid) if _map_headers(r)), None)
    if header_idx is not None:
        mapping = _map_headers(grid[header_idx])
        for n, cells in enumerate(grid[header_idx + 1:], start=1):
            if _map_headers(cells):
                continue
            raw = {f: (cells[i] if i < len(cells) else "") for i, f in mapping.items()}
            _normalize_row(raw, n, result)
        return result

    # Layout 2: labelled blocks. A new block starts when a field repeats.
    state = {"cur": {}, "last": None, "n": 0}

    def flush():
        if state["cur"]:
            state["n"] += 1
            _normalize_row(state["cur"], state["n"], result)
        state["cur"], state["last"] = {}, None

    for line in lines:
        if PAGE_NOISE_RE.match(line):
            continue
        m = LABEL_RE.match(line)
        key = _label_key(m.group(1)) if m else None
        # a bare "A) ..." only counts as an option once the question has started
        if key and key in OPTION_KEYS and "question" not in state["cur"]:
            key = None
        if key:
            if key in state["cur"]:          # field repeated -> a new question block begins
                flush()
            state["cur"][key] = m.group(2)
            state["last"] = key
        elif line.strip() and state["last"] in TEXT_KEYS:
            state["cur"][state["last"]] += " " + line.strip()   # wrapped text
    flush()

    if not result.rows and not result.errors:
        result.errors.append({
            "row": 1,
            "reason": "No questions found. Use a table with Subject, Difficulty, Topic, Question, "
                      "Option A-D, Correct Answer and Solution columns, or blocks starting with 'Subject:'.",
        })
    return result


def parse_file(filename: str, data: bytes) -> ParseResult:
    name = filename.lower()
    if name.endswith(".csv"):
        return parse_csv(data)
    if name.endswith((".xlsx", ".xlsm")):
        return parse_xlsx(data)
    if name.endswith(".pdf"):
        return parse_pdf(data)
    raise ValueError("Unsupported file type. Upload a .csv, .xlsx or .pdf file.")
